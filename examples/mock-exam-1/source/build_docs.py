"""Build the exam paper and answer key (.docx) from exam_content.py."""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

import exam_content as C

BLACK = RGBColor(0, 0, 0)
CODE = RGBColor(0xC2, 0x56, 0x0B)
GREY = RGBColor(0x55, 0x55, 0x55)
BODY_FONT = "Calibri"
MONO = "Courier New"
PAGE_W_CM = 21.0 - 2 * 2.0          # A4 minus 2 cm margins
DXA_PER_CM = 567


def new_doc():
    d = Document()
    for s in d.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.space_after = Pt(4)
    return d


def add_runs(p, text, size=None, color=None, bold=False):
    """Plain text with `inline code` spans rendered in Courier New / C2560B."""
    for i, part in enumerate(re.split(r"`([^`]*)`", text)):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if i % 2:
            r.font.name = MONO
            r.font.color.rgb = CODE
            r.font.size = Pt((size or 10.5) - 0.5)
        elif color is not None:
            r.font.color.rgb = color
    return p


def heading(d, text, level=1):
    sizes = {0: 20, 1: 14, 2: 12, 3: 11}
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level <= 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(sizes[level])
    r.font.color.rgb = BLACK
    return p


def rule(d):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "E0E0E0")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _width(table, widths_cm):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths_cm) * DXA_PER_CM)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)
            tcW = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tcW.set(qn("w:w"), str(int(w * DXA_PER_CM)))
            tcW.set(qn("w:type"), "dxa")


def _borders(table, color="D0D0D0"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", color)):
            e.set(qn(k), v)
        b.append(e)
    tblPr.append(b)


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    trPr.append(c)


def code_block(d, lines, fill="F5F5F5"):
    t = d.add_table(rows=1, cols=1)
    _borders(t, "E0E0E0")
    _width(t, [PAGE_W_CM])
    _cant_split(t.rows[0])
    cell = t.rows[0].cells[0]
    _shade(cell, fill)
    cell.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.font.name = MONO
        r.font.size = Pt(9)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


def callout(d, title, lines, fill="FFF4E5"):
    t = d.add_table(rows=1, cols=1)
    _borders(t, "F0C080")
    _width(t, [PAGE_W_CM])
    _cant_split(t.rows[0])
    cell = t.rows[0].cells[0]
    _shade(cell, fill)
    p = cell.paragraphs[0]
    add_runs(p, title, bold=True)
    for line in lines:
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(1)
        add_runs(q, line)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


def grid(d, header, rows, widths, header_fill="F0F0F0", size=9.5):
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t)
    _width(t, widths)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        _shade(c, header_fill)
        add_runs(c.paragraphs[0], h, size=size, bold=True)
    for ri, row in enumerate(rows, start=1):
        _cant_split(t.rows[ri])
        for ci, val in enumerate(row):
            add_runs(t.rows[ri].cells[ci].paragraphs[0], str(val), size=size)
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def numbered(d, items):
    for i, it in enumerate(items, 1):
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        add_runs(p, f"{i}.  {it}")


def bullets(d, items, indent=0.6):
    for it in items:
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent + 0.5)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, "–  " + it)


def title_block(d, title, sub):
    p = d.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = BLACK
    p = d.add_paragraph()
    r = p.add_run(sub)
    r.font.size = Pt(12)
    r.font.color.rgb = GREY
    rule(d)


def environment(d):
    heading(d, "Environment", 1)
    grid(d, ["Node", "Role", "Hostname", "IPv4 (rhcsa-exam network)", "Spare disks", "Software"],
         [list(e) for e in C.ENVIRONMENT], [1.2, 2.3, 3.5, 2.4, 3.3, 4.3], size=9)


# ----------------------------------------------------------------------------- paper
def build_paper(path):
    d = new_doc()
    title_block(d, C.TITLE, f"{C.SUBTITLE}  ·  two nodes  ·  RHEL 9  ·  {C.MINUTES} minutes")
    environment(d)
    heading(d, "Before you start the clock", 1)
    numbered(d, C.SETUP_STEPS)
    code_block(d, C.SETUP_CODE)
    p = d.add_paragraph()
    add_runs(p, "Snapshots are your reset button: to sit the exam again, restore both snapshots above and repeat this setup.")
    heading(d, "Rules", 1)
    bullets(d, C.RULES, indent=0)
    rule(d)
    heading(d, "Tasks", 1)
    for t in C.TASKS:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = bool(t.get("bullets"))
        r = p.add_run(f"{t['n']}.  ")
        r.bold = True
        r2 = p.add_run(f"On {t['node']}: ")
        r2.bold = True
        add_runs(p, t["text"])
        if t.get("bullets"):
            bullets(d, t["bullets"])
    rule(d)
    heading(d, "When time is up", 1)
    numbered(d, [C.GRADING_STEPS[0]])
    p = d.add_paragraph(); add_runs(p, "2.  " + C.GRADING_STEPS[1])
    code_block(d, C.GRADING_CODE_1)
    p = d.add_paragraph(); add_runs(p, "3.  " + C.GRADING_STEPS[2])
    code_block(d, C.GRADING_CODE_2)
    p = d.add_paragraph()
    add_runs(p, "The grader is read-only. It prints PASS/FAIL per task with the check that failed, and a per-objective score like Red Hat's report.", color=GREY)
    d.save(path)


# ----------------------------------------------------------------------------- key
def build_key(path):
    d = new_doc()
    title_block(d, f"{C.TITLE} — Answer Key", f"{C.SUBTITLE}  ·  do not read before sitting the exam")
    heading(d, "Dependency map", 1)
    p = d.add_paragraph()
    add_runs(p, "Tasks stack, as on the real exam. A failure on the left breaks the tasks on the right. The grader prints a note when that happens.")
    dep_rows = []
    for t in C.TASKS:
        users = [u["n"] for u in C.TASKS if t["n"] in u["deps"]]
        if users:
            dep_rows.append([f"{t['n']} ({t['node']})", ", ".join(str(u) for u in users)])
    grid(d, ["Task", "Required by"], dep_rows, [3.0, 14.0])
    heading(d, "How this paper was verified", 1)
    bullets(d, [
        "The commands in this key were run on a two-node RHEL 9.0 lab (VirtualBox, clean snapshots), both nodes were rebooted, and the grader scored 300/300 on each.",
        "On the untouched lab the grader scores 0. Not one task passes by default.",
        "Three common wrong answers were planted and the grader failed each one: ftp allowed only in the runtime firewall config, a /webfiles label without a semanage fcontext rule, and autofs started but not enabled.",
    ], indent=0)
    heading(d, "Objective coverage", 1)
    cats = {}
    for t in C.TASKS:
        cats.setdefault(t["cat"], []).append(str(t["n"]))
    grid(d, ["Objective (score report line)", "Tasks"],
         [[C.CAT_NAMES[c], ", ".join(v)] for c, v in cats.items()], [10.0, 7.0])
    for t in C.TASKS:
        rule(d)
        heading(d, f"Task {t['n']}  ·  {t['node']}  ·  {C.CAT_NAMES[t['cat']]}", 2)
        p = d.add_paragraph()
        add_runs(p, t["text"], color=GREY)
        if t.get("bullets"):
            for b in t["bullets"]:
                q = d.add_paragraph()
                q.paragraph_format.left_indent = Cm(0.8)
                q.paragraph_format.space_after = Pt(1)
                add_runs(q, "–  " + b, color=GREY)
        if t["deps"]:
            p = d.add_paragraph()
            add_runs(p, "Depends on: " + ", ".join(f"task {x}" for x in t["deps"]), bold=True)
        heading(d, "Solution", 3)
        code_block(d, t["steps"])
        heading(d, "Prove it", 3)
        code_block(d, t["verify"], fill="EEF5EE")
        p = d.add_paragraph()
        add_runs(p, "Expected: ", bold=True)
        add_runs(p, t["expect"])
        if t["alt"]:
            heading(d, "Also correct", 3)
            bullets(d, t["alt"], indent=0)
        if t["traps"]:
            heading(d, "Wrong answers that look right", 3)
            bullets(d, t["traps"], indent=0)
    rule(d)
    heading(d, "Grading your attempt", 1)
    numbered(d, [C.GRADING_STEPS[0]])
    p = d.add_paragraph(); add_runs(p, "2.  " + C.GRADING_STEPS[1])
    code_block(d, C.GRADING_CODE_1)
    p = d.add_paragraph(); add_runs(p, "3.  " + C.GRADING_STEPS[2])
    code_block(d, C.GRADING_CODE_2)
    p = d.add_paragraph()
    add_runs(p, "Paste both outputs to Claude (with the rhcsa-mock-exam skill) to have the result explained task by task.")
    d.save(path)


if __name__ == "__main__":
    out = sys.argv[1]
    build_paper(f"{out}/RHCSA9_Mock_Exam_1.docx")
    build_key(f"{out}/RHCSA9_Mock_Exam_1_Key.docx")
    print("built")
